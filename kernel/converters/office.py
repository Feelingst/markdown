from __future__ import annotations

from .utils import clean_markdown, heading, stem_title, table_from_rows


def from_pdf(data: bytes, filename: str) -> str:
    import io

    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts = [heading(1, stem_title(filename))]
    parts.append(f"_Páginas: {len(reader.pages)}_\n\n")

    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        parts.append(heading(2, f"Página {index}"))
        parts.append(text + "\n\n" if text else "_Sin texto extraíble en esta página._\n\n")

    return clean_markdown("".join(parts))


def from_docx(data: bytes, filename: str) -> str:
    import io

    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(io.BytesIO(data))
    parts: list[str] = []
    saw_title = False

    def iter_block_items(parent):
        parent_elm = parent.element.body
        for child in parent_elm.iterchildren():
            if child.tag.endswith("}p"):
                yield Paragraph(child, parent)
            elif child.tag.endswith("}tbl"):
                yield Table(child, parent)

    for block in iter_block_items(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            style = (block.style.name or "").lower() if block.style else ""
            if "heading 1" in style or style == "title":
                parts.append(heading(1, text))
                saw_title = True
            elif "heading 2" in style:
                parts.append(heading(2, text))
            elif "heading 3" in style:
                parts.append(heading(3, text))
            elif "heading 4" in style:
                parts.append(heading(4, text))
            elif "list" in style:
                parts.append(f"- {text}\n")
            else:
                parts.append(f"{text}\n\n")
        else:
            rows = []
            for row in block.rows:
                rows.append([cell.text.strip() for cell in row.cells])
            parts.append(table_from_rows(rows))

    if not parts:
        return clean_markdown(f"{heading(1, stem_title(filename))}_Documento vacío._\n")
    if not saw_title:
        parts.insert(0, heading(1, stem_title(filename)))
    return clean_markdown("".join(parts))


def from_pptx(data: bytes, filename: str) -> str:
    import io

    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    parts = [heading(1, stem_title(filename))]

    for index, slide in enumerate(presentation.slides, start=1):
        parts.append(heading(2, f"Diapositiva {index}"))
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
            if shape.has_table:
                rows = []
                for row in shape.table.rows:
                    rows.append([cell.text.strip() for cell in row.cells])
                parts.append(table_from_rows(rows))
        if texts:
            parts.append("\n\n".join(texts) + "\n\n")
        else:
            parts.append("_Sin texto en esta diapositiva._\n\n")

    return clean_markdown("".join(parts))


def from_xlsx(data: bytes, filename: str) -> str:
    import io

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts = [heading(1, stem_title(filename))]

    for sheet in workbook.worksheets:
        parts.append(heading(2, sheet.title))
        rows = []
        for row in sheet.iter_rows(values_only=True):
            if all(cell is None or str(cell).strip() == "" for cell in row):
                continue
            rows.append(["" if cell is None else str(cell) for cell in row])
        parts.append(table_from_rows(rows) if rows else "_Hoja vacía._\n\n")

    workbook.close()
    return clean_markdown("".join(parts))
